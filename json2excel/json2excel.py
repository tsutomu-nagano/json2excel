import yaml
import openpyxl
import pandas as pd
import re
import json
import sys
import copy
import shutil
import utils
from typing import List
from excel_decorator import ExcelDecorator

class JSON2Excel():

    __template_sheet_names : List[str] = []

    @staticmethod
    def rename_header(df, colmaps):

        # 文字列一致で置換
        df = df.rename(columns = {colmap["var"]:colmap["ja"]["text"] for colmap in colmaps if "text" in colmap["ja"]})

        # # 正規表現で置換
        # headers  = df.columns.values
        # patterns = {colmap["ja"]["re"]:colmap["var"] for colmap in colmaps if "re" in colmap["ja"]}
        # if len(patterns) >= 1:
        #     for pattern, var in patterns.items():
                
        #         headers = [re.sub(pattern, var, h) for h in headers]
            
        #     df.columns = headers

        return(df)







    @classmethod
    def json_with_yaml2xls_core(cls, y, j, wb, sheet = ""):
        if isinstance(y, dict):
            if "type" in y:
                conv_type = y["type"]

                if sheet != "":
                    y["sheet"] = sheet

                if conv_type == "cell":
                    sheet = y["sheet"]
                    address = y["address"]
                    # wb = openpyxl.load_workbook(dest)
                    wb[sheet][address].value = j
                    # wb.save(dest)
                    # wb.close()

                
                if conv_type == "table":
                    sheet = y["sheet"]
                    address = y["address"]
                    r = utils.a1_to_range(address)[0]
                    colmaps = y["colmap"]

                    df = pd.DataFrame(j)
                    
                    # 名前変換
                    df = JSON2Excel.rename_header(df, colmaps)

                    ws = wb[sheet]
                    for idx, value in enumerate(df.columns.values):
                        ws.cell(row = r["row"], column= r["col"] + idx, value = value)
                    # with pd.ExcelWriter(wb, engine='openpyxl', mode="a", if_sheet_exists='overlay') as writer:
                    #     pd.DataFrame([df.columns.values]).to_excel(writer, sheet_name=sheet, startrow=r["row"] - 1, startcol=r["col"] - 1, index=False, header = False)


                    for i, row in df.iterrows():
                        for j, value in enumerate(row):
                            ws.cell(row=r["row"] + 1 + i, column=r["col"] + j, value=value)

                    # with pd.ExcelWriter(wb, engine='openpyxl', mode="a", if_sheet_exists='overlay') as writer:
                    #     df.to_excel(writer, sheet_name=sheet, startrow=r["row"], startcol=r["col"] - 1, index=False, header = False)


                if conv_type == "list":

                    m = re.match("^\\$(?P<name>.+)@(?P<template>.+)", y["sheet"])
                    sheet_name_from = m.group("name")
                    template_sheet_name = m.group("template")

                    cls.__template_sheet_names.append(template_sheet_name)

                    # wb = openpyxl.load_workbook(dest)
                    ws_temp = wb[template_sheet_name]
                    for j_ in j:
                        ws_copy = wb.copy_worksheet(ws_temp)
                        ws_copy.title = j_[sheet_name_from]

                    # wb.save(dest)
                    # wb.close()

                    for j_ in j:
                        sheet = j_[sheet_name_from]
                        JSON2Excel.json_with_yaml2xls_core(copy.deepcopy(y["listitem"]), j_ , wb, sheet) 


            else:
                for k in y:
                    JSON2Excel.json_with_yaml2xls_core(y[k], j[k], wb, sheet)


    @classmethod
    def json_with_yaml2xls(cls, src, config, temp, dest):
        shutil.copy(temp, dest)
        with open(config, 'r') as ymlf, \
             open(src, 'r') as jsonf:

            wb = openpyxl.load_workbook(dest)


            cls.json_with_yaml2xls_core(
                y = yaml.safe_load(ymlf),
                j = json.load(jsonf),
                wb = wb
            )

            [wb.remove(wb[s]) for s in cls.__template_sheet_names]

            wb.save(dest)
            wb.close()


if __name__ == "__main__":
    args = sys.argv

    src = args[1]
    config = args[2]
    dest = args[3]
    temp = args[4]

    JSON2Excel.json_with_yaml2xls(src, config, temp, dest)
